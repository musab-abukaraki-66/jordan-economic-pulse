"""M02 recon: profile every raw file in data/raw (read-only)."""
import zipfile, io, glob, os, sys
from pathlib import Path
REPO = Path(__file__).resolve().parents[2]
import pandas as pd
import openpyxl

ROOT = str(REPO)
OUT = os.path.join(ROOT, "docs", "audit")
os.makedirs(OUT, exist_ok=True)

# ---------- World Bank ----------
rows = []
for z in sorted(glob.glob(os.path.join(ROOT, "data", "raw", "worldbank", "*.zip"))):
    with zipfile.ZipFile(z) as zf:
        names = zf.namelist()
        api = [n for n in names if n.startswith("API_")][0]
        df = pd.read_csv(zf.open(api), skiprows=4)
    df = df.loc[:, ~df.columns.str.startswith("Unnamed")]
    years = [c for c in df.columns if c.isdigit()]
    jo = df[df["Country Code"] == "JOR"]
    code = df["Indicator Code"].iloc[0]
    name = df["Indicator Name"].iloc[0]
    if len(jo):
        s = jo[years].iloc[0]
        nn = s.dropna()
        first = nn.index.min() if len(nn) else None
        last = nn.index.max() if len(nn) else None
        n_nn = len(nn)
        lastval = nn.iloc[-1] if len(nn) else None
    else:
        first = last = None; n_nn = 0; lastval = None
    rows.append(dict(file=os.path.basename(z), code=code, name=name, n_rows=len(df),
                     n_countries=df["Country Code"].nunique(), yr_min=years[0], yr_max=years[-1],
                     jor_rows=len(jo), jor_first=first, jor_last=last, jor_nonnull=n_nn, jor_lastval=lastval,
                     zip_members=";".join(names)))
wb = pd.DataFrame(rows)
wb.to_csv(os.path.join(OUT, "recon_worldbank.csv"), index=False)
pd.set_option("display.width", 250); pd.set_option("display.max_colwidth", 60); pd.set_option("display.max_rows", 200)
print(wb.drop(columns=["file", "zip_members"]).to_string())

# ---------- CBJ ----------
print("\n===== CBJ =====")
for f in sorted(glob.glob(os.path.join(ROOT, "data", "raw", "cbj", "*.xlsx"))):
    wbk = openpyxl.load_workbook(f, data_only=True)
    print(f"\n#### {os.path.basename(f)}  sheets={wbk.sheetnames}")
    for ws in wbk.worksheets:
        print(f"  -- sheet '{ws.title}' dims={ws.dimensions} max_row={ws.max_row} max_col={ws.max_column} merged={len(ws.merged_cells.ranges)}")
        for r in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 14), values_only=True):
            vals = [("" if v is None else str(v)[:22]) for v in r[:14]]
            if any(vals):
                print("   |", " | ".join(vals))
