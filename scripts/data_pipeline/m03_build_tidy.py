"""M03: parse raw WB zips + CBJ bulletin sheets into tidy long staging tables and audit them.
Raw files are read-only. Outputs: data/staging/*.csv, docs/audit/m03_audit.txt
"""
import glob, os, re, zipfile, io, json
from pathlib import Path
REPO = Path(__file__).resolve().parents[2]
import pandas as pd
import openpyxl

ROOT = str(REPO)
RAW_WB = os.path.join(ROOT, "data", "raw", "worldbank")
RAW_CBJ = os.path.join(ROOT, "data", "raw", "cbj")
STG = os.path.join(ROOT, "data", "staging")
DOC = os.path.join(ROOT, "docs", "audit")
os.makedirs(STG, exist_ok=True); os.makedirs(DOC, exist_ok=True)

LOG = []
EXCL = []
def log(*a):
    s = " ".join(str(x) for x in a); LOG.append(s); print(s)

YR = re.compile(r"^(\(\d\))?\s*((?:19|20)\d\d)$")
def parse_year(v):
    """returns (year, footnote_flag) or (None, False)"""
    if v is None: return None, False
    if isinstance(v, float) and v.is_integer(): v = int(v)
    s = str(v).strip()
    m = YR.match(s)
    if not m: return None, False
    return int(m.group(2)), bool(m.group(1))

# ---------------------------------------------------------------- World Bank
PEERS = ["JOR", "EGY", "LBN", "MAR", "TUN", "SAU", "ARE", "TUR", "IRQ", "MEA", "LMC", "WLD"]
wb_rows = []; wb_meta = []
for z in sorted(glob.glob(os.path.join(RAW_WB, "*.zip"))):
    with zipfile.ZipFile(z) as zf:
        api = [n for n in zf.namelist() if n.startswith("API_")][0]
        df = pd.read_csv(zf.open(api), skiprows=4)
    years = [c for c in df.columns if c.isdigit()]
    code = df["Indicator Code"].iloc[0]; name = df["Indicator Name"].iloc[0]
    wb_meta.append((code, name, os.path.basename(z)))
    sub = df[df["Country Code"].isin(PEERS)]
    long = sub.melt(id_vars=["Country Name", "Country Code"], value_vars=years, var_name="Year", value_name="Value")
    long["Year"] = long["Year"].astype(int)
    long["IndicatorCode"] = code
    wb_rows.append(long)
wb = pd.concat(wb_rows, ignore_index=True)
log("WB long rows (peer set, incl. nulls):", len(wb))
wb_nn = wb.dropna(subset=["Value"])
log("WB non-null rows:", len(wb_nn))
dup = wb_nn.duplicated(["Country Code", "IndicatorCode", "Year"]).sum()
log("WB duplicate (country,indicator,year):", dup)
wb_nn.to_csv(os.path.join(STG, "wb_long.csv"), index=False)
pd.DataFrame(wb_meta, columns=["IndicatorCode", "IndicatorName", "SourceFile"]).to_csv(os.path.join(STG, "wb_indicators.csv"), index=False)

# ---------------------------------------------------------------- CBJ trade tables
def col(ws, letter, row): return ws[f"{letter}{row}"].value

TABLES = {
  # name: (file, ycol, {col: (item, group)}, total_col, unit, flow)
  "exports_commodity": ("domestic_exports_by_commodity_according_to_s.i.t.c..xlsx", "O",
      {"C":"Other","D":"Misc. Manufactured Articles","E":"Machinery & Transport Equipment","F":"Manufactured Goods (by Material)",
       "G":"Chemicals","H":"Animal & Vegetable Oils & Fats","I":"Mineral Fuels & Lubricants","J":"Crude Materials (Inedible)",
       "K":"Beverages & Tobacco","L":"Food & Live Animals"}, "M", "JD thousand", "Domestic exports"),
  "imports_commodity": ("imports_by_commodity_according_to_s.i.t.c..xlsx", "O",
      {"C":"Other","D":"Misc. Manufactured Articles","E":"Machinery & Transport Equipment","F":"Manufactured Goods (by Material)",
       "G":"Chemicals","H":"Animal & Vegetable Oils & Fats","I":"Mineral Fuels & Lubricants","J":"Crude Materials (Inedible)",
       "K":"Beverages & Tobacco","L":"Food & Live Animals"}, "M", "JD thousand", "Imports"),
  "exports_function": ("domestic_exports_by_economic_function.xlsx", "K",
      {"C":"Other Goods n.e.c.","D":"Capital Goods","E":"Crude & Intermediate - Other","F":"Construction Materials",
       "G":"Durable Consumer Goods","H":"Current Consumer Goods"}, "I", "JD thousand", "Domestic exports"),
  "imports_function": ("imports_by_economic_function.xlsx", "M",
      {"C":"Other Goods n.e.c.","D":"Capital - Other","E":"Capital - Other Machinery & Equipment","F":"Capital - Machinery & Transport Equipment",
       "G":"Crude & Intermediate - Other","H":"Crude & Intermediate - Oil & Fuels","I":"Durable Consumer Goods","J":"Current Consumer Goods"},
      "K", "JD thousand", "Imports"),
  "exports_geo": ("geographic_distribution_of_exports.xlsx", "M",
      {"C":"Other Countries","D":"Japan","E":"India","F":"China","G":"U.S.A.","H":"Other European Countries",
       "I":"European Union","J":"Arab Countries"}, "K", "JD thousand", "Exports"),
  "imports_geo": ("geographic_distribution_of_imports.xlsx", "M",
      {"C":"Other Countries","D":"Japan","E":"India","F":"China","G":"U.S.A.","H":"Other European Countries",
       "I":"European Union","J":"Arab Countries"}, "K", "JD thousand", "Imports"),
  "imports_arab_commodity": ("imports_from_arab_countries_according_to_s.i.t.c..xlsx", "N",
      {"B":"Other","C":"Misc. Manufactured Articles","D":"Machinery & Transport Equipment","E":"Manufactured Goods (by Material)",
       "F":"Chemicals","G":"Animal & Vegetable Oils & Fats","H":"Mineral Fuels & Lubricants","I":"Crude Materials (Inedible)",
       "J":"Beverages & Tobacco","K":"Food & Live Animals"}, "L", "JD thousand", "Imports from Arab countries"),
}
INDICES = {
  "idx_price_exports": ("index_of_unit_price_of_domestic_exports.xlsx", "O",
      {"C":"Misc. Manufactured Articles","G":"Manufactured Goods (by Material)","H":"Chemicals","J":"Crude Materials (Inedible)",
       "K":"Beverages & Tobacco","L":"Food & Live Animals"}, "M", "Unit price index", "Domestic exports"),
  "idx_volume_exports": ("index_of_unit_volume_of_domestic_exports.xlsx", "O",
      {"C":"Misc. Manufactured Articles","G":"Manufactured Goods (by Material)","H":"Chemicals","J":"Crude Materials (Inedible)",
       "K":"Beverages & Tobacco","L":"Food & Live Animals"}, "M", "Unit volume index", "Domestic exports"),
  "idx_price_imports": ("index_of_unit_price_of_imports_..xlsx", "Q",
      {"D":"Misc. Manufactured Articles","F":"Machinery & Transport Equipment","G":"Manufactured Goods (by Material)","H":"Chemicals",
       "I":"Animal & Vegetable Oils & Fats","J":"Mineral Fuels & Lubricants","K":"Crude Materials (Inedible)",
       "L":"Beverages & Tobacco","M":"Food & Live Animals"}, "N", "Unit price index", "Imports"),
  "idx_volume_imports": ("index_of_unit_volume_of_imports.xlsx", "O",
      {"B":"Misc. Manufactured Articles","E":"Machinery & Transport Equipment","F":"Manufactured Goods (by Material)","G":"Chemicals",
       "H":"Animal & Vegetable Oils & Fats","I":"Mineral Fuels & Lubricants","J":"Crude Materials (Inedible)",
       "K":"Beverages & Tobacco","L":"Food & Live Animals"}, "M", "Unit volume index", "Imports"),
}

def read_table(fname, ycol, cols, totcol):
    ws = openpyxl.load_workbook(os.path.join(RAW_CBJ, fname), data_only=True).worksheets[0]
    recs = []; skipped = []
    for r in range(1, ws.max_row + 1):
        y, fn = parse_year(col(ws, ycol, r))
        if y is None: continue
        nums = [col(ws, c, r) for c in list(cols) + ([totcol] if totcol else [])]
        if not any(isinstance(v, (int, float)) for v in nums):
            skipped.append((r, y)); continue
        rec = {"Year": y, "YearFootnote": fn, "SrcRow": r}
        for c, item in cols.items(): rec[item] = col(ws, c, r)
        if totcol: rec["__TOTAL__"] = col(ws, totcol, r)
        recs.append(rec)
    out = pd.DataFrame(recs)
    dupmask = out.duplicated("Year", keep="first")
    if dupmask.any():
        for _, rr in out[dupmask].iterrows():
            first = out[(out.Year == rr.Year) & ~dupmask].iloc[0]
            same = all(rr[k] == first[k] for k in out.columns if k not in ("SrcRow", "YearFootnote"))
            EXCL.append((fname, int(rr.SrcRow), int(rr.Year), "duplicate year label; values %s to first occurrence (row %d) - excluded, next year missing in source" % ("identical" if same else "DIFFER", first.SrcRow)))
        out = out[~dupmask].reset_index(drop=True)
    return out, skipped

audit = {}
long_frames = []
for name, (fname, ycol, cols, totcol, unit, flow) in TABLES.items():
    df, skipped = read_table(fname, ycol, cols, totcol)
    log(f"\n=== {name}: {fname}")
    log(f" rows={len(df)} years {df.Year.min()}..{df.Year.max()} distinct years={df.Year.nunique()} dup years={df.Year.duplicated().sum()} skipped_no_values={skipped}")
    allyrs = set(range(df.Year.min(), df.Year.max() + 1)); missing = sorted(allyrs - set(df.Year))
    log(f" missing years in range: {missing}")
    items = list(cols.values())
    for it in items: df[it] = pd.to_numeric(df[it], errors="coerce")
    df["__TOTAL__"] = pd.to_numeric(df["__TOTAL__"], errors="coerce")
    comp_sum = df[items].sum(axis=1, min_count=1)
    diff = (comp_sum - df["__TOTAL__"])
    rel = (diff.abs() / df["__TOTAL__"].abs().where(df["__TOTAL__"].abs() > 0))
    bad = df.assign(comp_sum=comp_sum, diff=diff).loc[(diff.abs() > 5) & (rel > 0.002), ["Year", "__TOTAL__", "comp_sum", "diff"]]
    log(f" total-vs-components: max|diff|={diff.abs().max():.2f}, rows with diff>5 & >0.2%: {len(bad)}")
    if len(bad): log(bad.head(12).to_string())
    log(f" null cells per item: { {it:int(df[it].isna().sum()) for it in items} }")
    log(f" zero cells per item: { {it:int((df[it]==0).sum()) for it in items} }")
    log(f" negative cells: { {it:int((df[it]<0).sum()) for it in items if (df[it]<0).any()} }")
    log(f" year footnote rows: {df.loc[df.YearFootnote,'Year'].tolist()}")
    l = df.melt(id_vars=["Year", "SrcRow"], value_vars=items, var_name="Category", value_name="ValueJDk")
    l["Table"] = name; l["Flow"] = flow
    l.to_csv(os.path.join(STG, f"cbj_{name}_long.csv"), index=False)
    tot = df[["Year", "__TOTAL__"]].rename(columns={"__TOTAL__": "GrandTotalJDk"}); tot["Table"] = name
    tot.to_csv(os.path.join(STG, f"cbj_{name}_total.csv"), index=False)
    dup_grain = l.duplicated(["Year", "Category"]).sum()
    log(f" GRAIN Year x Category -> long rows={len(l)}, duplicate keys={dup_grain}")

for name, (fname, ycol, cols, totcol, unit, flow) in INDICES.items():
    df, skipped = read_table(fname, ycol, cols, totcol)
    log(f"\n=== {name}: {fname}")
    log(f" rows={len(df)} years {df.Year.min()}..{df.Year.max()} dup years={df.Year.duplicated().sum()} skipped={skipped}")
    allyrs = set(range(df.Year.min(), df.Year.max() + 1)); log(f" missing years: {sorted(allyrs - set(df.Year))}")
    items = list(cols.values())
    for it in items + ["__TOTAL__"]: df[it] = pd.to_numeric(df[it], errors="coerce")
    log(f" null cells: { {it:int(df[it].isna().sum()) for it in items+['__TOTAL__']} }")
    log(f" value range: min={df[items+['__TOTAL__']].min().min():.1f} max={df[items+['__TOTAL__']].max().max():.1f}")
    l = df.melt(id_vars=["Year", "SrcRow"], value_vars=items + ["__TOTAL__"], var_name="Category", value_name="IndexValue")
    l["Category"] = l["Category"].replace({"__TOTAL__": "General Index"}); l["Table"] = name; l["Flow"] = flow
    l.to_csv(os.path.join(STG, f"cbj_{name}_long.csv"), index=False)
    log(f" GRAIN Year x Category dup keys={l.duplicated(['Year','Category']).sum()}")

# ---------------------------------------------------------------- Monetary survey
ws = openpyxl.load_workbook(os.path.join(RAW_CBJ, "3__monetary__survey_of_the_banking_system.xlsx"), data_only=True).worksheets[0]
old_cols = {"C":"Other Items (Net)","D":"Capital, Reserves & Allowances","E":"Foreign Liabilities","F":"Government Deposits",
            "G":"Quasi-Money","H":"Money Supply (M1)","J":"Claims on Government","K":"Claims on Municipalities & Public Entities",
            "L":"Claims on Private Sector","M":"Foreign Assets"}
new_cols = {"C":"Money Supply (M2)","E":"Other Items (Net)","F":"Claims on Financial Institutions","G":"Claims on Private Sector",
            "I":"Claims on Public Entities","J":"Net Claims on Central Government (Own Budget)","K":"Net Claims on Central Government (General Budget)",
            "M":"Foreign Assets - Licensed Banks","N":"Foreign Assets - Central Bank","L":"Foreign Assets (Total)","D":"Domestic Assets (Total)","H":"Net Claims on Public Sector (Total)"}
mon = []
for r in range(1, ws.max_row + 1):
    if r <= 45:  # old table block
        y, fn = parse_year(col(ws, "N", r))
        if y and isinstance(col(ws, "C", r), (int, float)):
            for c, it in old_cols.items(): mon.append((y, "1964-1992 basis", it, col(ws, c, r), fn))
            mon.append((y, "1964-1992 basis", "Total Assets = Liabilities", col(ws, "I", r), fn))
    elif 66 <= r <= 96:
        y, fn = parse_year(col(ws, "O", r))
        if y and isinstance(col(ws, "C", r), (int, float)):
            for c, it in new_cols.items(): mon.append((y, "1993-2021 basis", it, col(ws, c, r), fn))
mon = pd.DataFrame(mon, columns=["Year", "Basis", "Item", "ValueJDm", "YearFootnote"])
# 1998 appears twice in the 1993-2021 block: plain row + '(7)1998' (public sector accounts reclassified). Keep the restated one.
sup = (mon.Basis == "1993-2021 basis") & (mon.Year == 1998) & (~mon.YearFootnote)
EXCL.append(("3__monetary__survey_of_the_banking_system.xlsx", 71, 1998, "pre-reclassification 1998 row superseded by restated row 72 '(7)1998' (note 7: public sector accounts reclassified); differences e.g. M2 6003.3 vs 6026.3"))
mon = mon[~sup].reset_index(drop=True)
log("\n=== monetary survey"); log(mon.groupby("Basis").agg(rows=("Year", "size"), y0=("Year", "min"), y1=("Year", "max"), yrs=("Year", "nunique")).to_string())
log(" duplicate keys Year x Basis x Item:", mon.duplicated(["Year", "Basis", "Item"]).sum())
for b, g in mon.groupby("Basis"):
    p = g.pivot(index="Year", columns="Item", values="ValueJDm")
    log(f" {b}: null cells={int(p.isna().sum().sum())}, missing yrs={sorted(set(range(p.index.min(), p.index.max()+1))-set(p.index))}")
p_old = mon[mon.Basis == "1964-1992 basis"].pivot(index="Year", columns="Item", values="ValueJDm")
liab = p_old[["Other Items (Net)", "Capital, Reserves & Allowances", "Foreign Liabilities", "Government Deposits", "Quasi-Money", "Money Supply (M1)"]].sum(axis=1)
asset = p_old[["Claims on Government", "Claims on Municipalities & Public Entities", "Claims on Private Sector", "Foreign Assets"]].sum(axis=1)
log(f" old basis: max|liab-total|={(liab-p_old['Total Assets = Liabilities']).abs().max():.2f}, max|assets-total|={(asset-p_old['Total Assets = Liabilities']).abs().max():.2f}")
p_new = mon[mon.Basis == "1993-2021 basis"].pivot(index="Year", columns="Item", values="ValueJDm")
log(f" new basis: max|Domestic+Foreign-M2|={((p_new['Domestic Assets (Total)']+p_new['Foreign Assets (Total)'])-p_new['Money Supply (M2)']).abs().max():.2f} (assets identity check incl. other items)")
mon.to_csv(os.path.join(STG, "cbj_monetary_long.csv"), index=False)

# ---------------------------------------------------------------- Reserves
ws = openpyxl.load_workbook(os.path.join(RAW_CBJ, "9__central_bank_of_jordan_foreign_reserves.xlsx"), data_only=True).worksheets[0]
ymap = {}
from openpyxl.utils import get_column_letter as L
for ci in range(9, 17): ymap[L(ci)] = parse_year(ws.cell(11, ci).value)[0]        # I..P (2021..2014)
for ci in range(17, 38): ymap[L(ci)] = parse_year(ws.cell(7, ci).value)[0]        # Q..AK (2013..1993)
ROWS = {19:("Gold (ounces)","Ounces","Asset"),20:("Gold (value)","JD million","Asset"),22:("SDRs","JD million","Asset"),
        23:("Cash, Balances & Deposits","JD million","Asset"),24:("Bonds & Treasury Bills","JD million","Asset"),
        26:("Loans arising from payment agreements","JD million","Asset"),27:("Subscription to international financial institutions","JD million","Asset"),
        29:("Total assets in gold & foreign currencies","JD million","Total"),
        32:("Licensed banks deposits","JD million","Liability"),33:("Non-resident deposits","JD million","Liability"),
        34:("Central government deposits","JD million","Liability"),35:("Public entities deposits","JD million","Liability"),
        36:("Reserve deposits","JD million","Liability"),37:("Other liabilities","JD million","Liability"),
        39:("Total liabilities in foreign currencies","JD million","Total"),
        42:("Gross foreign reserves (JD)","JD million","Headline"),43:("Gross foreign reserves (USD)","USD million","Headline"),
        44:("Import coverage (months)","Months","Ratio")}
res = []
for r, (item, unit, kind) in ROWS.items():
    for c, y in ymap.items():
        v = ws[f"{c}{r}"].value
        res.append((y, item, unit, kind, v if isinstance(v, (int, float)) else None))
res = pd.DataFrame(res, columns=["Year", "Item", "Unit", "Kind", "Value"])
log("\n=== reserves"); log(f" years {res.Year.min()}..{res.Year.max()} distinct={res.Year.nunique()}, rows={len(res)}, dup keys={res.duplicated(['Year','Item']).sum()}")
nn = res.groupby("Item").Value.apply(lambda s: int(s.isna().sum()))
log(" null cells per item:", nn[nn > 0].to_dict())
pv = res.pivot(index="Year", columns="Item", values="Value")
assets = pv[["Gold (value)","SDRs","Cash, Balances & Deposits","Bonds & Treasury Bills","Loans arising from payment agreements","Subscription to international financial institutions"]].sum(axis=1)
log(f" assets identity max|sum-total|={(assets-pv['Total assets in gold & foreign currencies']).abs().max():.2f}")
liabs = pv[["Licensed banks deposits","Non-resident deposits","Central government deposits","Public entities deposits","Reserve deposits","Other liabilities"]].sum(axis=1)
log(f" liabilities identity max|sum-total|={(liabs-pv['Total liabilities in foreign currencies']).abs().max():.2f}")
gross = pv["Total assets in gold & foreign currencies"] - pv["Total liabilities in foreign currencies"]
log(f" gross reserves identity max|(assets-liab)-reported|={(gross-pv['Gross foreign reserves (JD)']).abs().max():.2f}")
res.to_csv(os.path.join(STG, "cbj_reserves_long.csv"), index=False)

# ---------------------------------------------------------------- WB vs CBJ reconciliation
jor = wb_nn[wb_nn["Country Code"] == "JOR"].pivot(index="Year", columns="IndicatorCode", values="Value")
ex = pd.read_csv(os.path.join(STG, "cbj_exports_geo_total.csv")).set_index("Year").GrandTotalJDk
im = pd.read_csv(os.path.join(STG, "cbj_imports_geo_total.csv")).set_index("Year").GrandTotalJDk
fx = jor["PA.NUS.FCRF"]
cmp = pd.DataFrame({"CBJ_exp_JDk": ex, "CBJ_imp_JDk": im, "FX": fx, "WB_exp_USD": jor["TX.VAL.MRCH.CD.WT"], "WB_imp_USD": jor["TM.VAL.MRCH.CD.WT"]}).dropna()
cmp["CBJ_exp_USD_at_WB_FX"] = cmp.CBJ_exp_JDk * 1000 / cmp.FX
cmp["CBJ_imp_USD_at_WB_FX"] = cmp.CBJ_imp_JDk * 1000 / cmp.FX
cmp["exp_ratio_WB_over_CBJ"] = cmp.WB_exp_USD / cmp.CBJ_exp_USD_at_WB_FX
cmp["imp_ratio_WB_over_CBJ"] = cmp.WB_imp_USD / cmp.CBJ_imp_USD_at_WB_FX
log("\n=== WB merchandise vs CBJ geographic totals (converted at WB official FX, 'LCU per USD')")
log(f" overlap years {cmp.index.min()}..{cmp.index.max()} n={len(cmp)}")
log(" export ratio WB/CBJ: min=%.3f median=%.3f max=%.3f" % (cmp.exp_ratio_WB_over_CBJ.min(), cmp.exp_ratio_WB_over_CBJ.median(), cmp.exp_ratio_WB_over_CBJ.max()))
log(" import ratio WB/CBJ: min=%.3f median=%.3f max=%.3f" % (cmp.imp_ratio_WB_over_CBJ.min(), cmp.imp_ratio_WB_over_CBJ.median(), cmp.imp_ratio_WB_over_CBJ.max()))
log(cmp.loc[[y for y in (1985, 1995, 2005, 2010, 2015, 2019, 2021) if y in cmp.index], ["CBJ_exp_JDk","FX","WB_exp_USD","exp_ratio_WB_over_CBJ","imp_ratio_WB_over_CBJ"]].round(3).to_string())
cmp.to_csv(os.path.join(DOC, "recon_wb_vs_cbj_trade.csv"))

# exports geo total == exports function total == exports commodity total ?
for a, b in (("exports_geo","exports_function"),("exports_geo","exports_commodity"),("imports_geo","imports_function"),("imports_geo","imports_commodity")):
    ta = pd.read_csv(os.path.join(STG, f"cbj_{a}_total.csv")).set_index("Year").GrandTotalJDk
    tb = pd.read_csv(os.path.join(STG, f"cbj_{b}_total.csv")).set_index("Year").GrandTotalJDk
    j = pd.concat([ta, tb], axis=1, keys=["a", "b"]).dropna()
    d = (j.a - j.b).abs()
    log(f" cross-table total {a} vs {b}: overlap={len(j)} yrs, exact-match yrs={(d<1).sum()}, max diff={d.max():.0f}, rel max={(d/j.a.abs()).max():.4f}")

pd.DataFrame(EXCL, columns=["File","SrcRow","Year","Reason"]).to_csv(os.path.join(DOC, "m03_excluded_rows.csv"), index=False)
log("\nEXCLUDED SOURCE ROWS:")
for e in EXCL: log(" ", e)
open(os.path.join(DOC, "m03_audit.txt"), "w", encoding="utf-8").write("\n".join(LOG))
